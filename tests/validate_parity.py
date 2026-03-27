"""validate_parity.py ? Compare Python engine outputs against workbook cached values.

Performs a two-source parity check:

  Source A (Python engine): calculate_portfolio() on workbook-loaded data
  Source B (Workbook cache): openpyxl data_only reads of Availability + CL sheets

Any metric that has a cached formula cell in the workbook is compared directly.
Metrics with no workbook cell (e.g., per-test excess amounts) are validated
against reference values from the implementation notes.

Usage:
    cd C:\\Users\\henry.yan\\borrowing-base
    PYTHONPATH=src python tests/validate_parity.py
    PYTHONPATH=src python tests/validate_parity.py path/to/workbook.xlsm
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from borrowing_base_workbench.calculator import calculate_portfolio
from borrowing_base_workbench.loader import load_workbook_data

try:
    from openpyxl import load_workbook as _xl_load
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

# ---------------------------------------------------------------------------
# Reference values (implementation notes, Phase 3+4, 2025-02-11 workbook)
# ---------------------------------------------------------------------------

REF = {
    # Asset counts
    "asset_count":                   41,
    "eligible_count":                25,
    "ineligible_count":              16,
    "vae_affected_count":            11,
    # Portfolio par
    "total_portfolio_par":           330_231_191,
    # Pre-concentration
    "total_pre_conc_eligible_value": 120_210_997,
    "total_borrowing_value":         73_085_307,
    # WAAR
    "raw_waar":                      0.6080,
    "discrete_obligor_count":        25,
    "waar_cap":                      0.65,
    "final_waar":                    0.6080,
    # Waterfall
    "total_excess":                  29_281_364,
    "net_abv":                       90_929_633,
    # Availability
    "test_a_facility":               200_000_000,
    "test_b_borrowing_base":         60_753_941,
    "test_c_credit_enhancement":     81_704_750,
    "availability":                  60_753_941,
    # Per-test excess
    "excess_max_obligors":           24_911_002,
    "excess_max_largest_industry":    4_370_362,
}

# Workbook cells to read from (data_only cached values)
# Sheet -> cell -> metric name
_AVAILABILITY_CELLS = {
    "L22": "wb_test_a",
    "L23": "wb_test_b",
    "L25": "wb_test_c",
    "L26": "wb_availability",
    "L48": "wb_waar",
}

_CL_CELLS = {
    "M43": "wb_total_abv",
    "M44": "wb_total_excess",
    "M45": "wb_net_abv",
}


# ---------------------------------------------------------------------------
# ANSI colour helpers
# ---------------------------------------------------------------------------

_GREEN  = "\033[92m"
_YELLOW = "\033[93m"
_RED    = "\033[91m"
_RESET  = "\033[0m"
_BOLD   = "\033[1m"


def _col(text: str, colour: str) -> str:
    try:
        # Only colour if stdout is a terminal
        if sys.stdout.isatty():
            return f"{colour}{text}{_RESET}"
    except AttributeError:
        pass
    return text


# ---------------------------------------------------------------------------
# Comparison helpers
# ---------------------------------------------------------------------------

def _dollars(v: float) -> str:
    if v is None:
        return "N/A"
    return f"${v:>16,.0f}"


def _pct(v: float) -> str:
    if v is None:
        return "N/A"
    return f"{v * 100:>12.4f}%"


def _count(v: int) -> str:
    return f"{v:>8d}"


def _cmp_dollars(engine: float, ref: float, tol: float = 100.0) -> str:
    diff = abs(engine - ref)
    if diff <= tol:
        return _col("PASS", _GREEN)
    elif diff <= 1000.0:
        return _col(f"WARN  D=${diff:,.0f}", _YELLOW)
    else:
        return _col(f"FAIL  D=${diff:,.0f}", _RED)


def _cmp_pct(engine: float, ref: float, tol: float = 0.0001) -> str:
    diff = abs(engine - ref)
    if diff <= tol:
        return _col("PASS", _GREEN)
    elif diff <= 0.001:
        return _col(f"WARN  D={diff:.4%}", _YELLOW)
    else:
        return _col(f"FAIL  D={diff:.4%}", _RED)


def _cmp_exact(engine, ref) -> str:
    if engine == ref:
        return _col("PASS", _GREEN)
    return _col(f"FAIL  got={engine!r}  exp={ref!r}", _RED)


def _row(label: str, engine_str: str, ref_str: str, status: str, source: str = "impl-note") -> str:
    return f"  {label:<42} {engine_str:>18}  {ref_str:>18}  {status}  [{source}]"


# ---------------------------------------------------------------------------
# Main validation routine
# ---------------------------------------------------------------------------

def run_parity(workbook_path: Path) -> int:
    """Run full parity check. Returns number of failures."""
    print(f"\n{'='*80}")
    print(f"  BORROWING BASE ENGINE ? PARITY VALIDATION")
    print(f"  Workbook: {workbook_path.name}")
    print(f"{'='*80}\n")

    # ---- Load data and run Python engine -----------------------------------
    print("  Loading workbook data...")
    data = load_workbook_data(workbook_path)
    print("  Running calculate_portfolio()...")
    calc = calculate_portfolio(data)
    w  = calc.waterfall
    av = w.availability if w else None

    failures = 0
    warnings = 0

    # ---- Read cached formula values from workbook --------------------------
    wb_values: dict[str, float | None] = {}
    if _HAS_OPENPYXL:
        try:
            wb = _xl_load(workbook_path, data_only=True, read_only=True)
            if "Availability" in wb.sheetnames:
                avail_ws = wb["Availability"]
                for cell, key in _AVAILABILITY_CELLS.items():
                    v = avail_ws[cell].value
                    wb_values[key] = float(v) if v is not None else None
            if "Concentration Limits" in wb.sheetnames:
                cl_ws = wb["Concentration Limits"]
                for cell, key in _CL_CELLS.items():
                    v = cl_ws[cell].value
                    wb_values[key] = float(v) if v is not None else None
            wb.close()
        except Exception as e:
            print(f"  Warning: could not read cached workbook cells: {e}")

    def _check_dollar(label: str, engine_val: float, ref_val: float,
                      wb_key: str | None = None, tol: float = 100.0):
        nonlocal failures, warnings
        # Compare engine vs reference
        status = _cmp_dollars(engine_val, ref_val, tol)
        if "FAIL" in status:
            failures += 1
        elif "WARN" in status:
            warnings += 1
        print(_row(label, _dollars(engine_val), _dollars(ref_val), status))

        # Additional: compare engine vs workbook cached cell
        if wb_key and wb_key in wb_values and wb_values[wb_key] is not None:
            wb_val = wb_values[wb_key]
            wb_status = _cmp_dollars(engine_val, wb_val, tol)
            if "FAIL" in wb_status:
                failures += 1
            elif "WARN" in wb_status:
                warnings += 1
            print(_row(f"  -> vs workbook {wb_key}", _dollars(engine_val),
                       _dollars(wb_val), wb_status, "workbook-cache"))

    def _check_rate(label: str, engine_val: float, ref_val: float,
                    wb_key: str | None = None, tol: float = 0.0001):
        nonlocal failures, warnings
        status = _cmp_pct(engine_val, ref_val, tol)
        if "FAIL" in status:
            failures += 1
        elif "WARN" in status:
            warnings += 1
        print(_row(label, _pct(engine_val), _pct(ref_val), status))
        if wb_key and wb_key in wb_values and wb_values[wb_key] is not None:
            wb_val = wb_values[wb_key]
            wb_status = _cmp_pct(engine_val, wb_val, tol)
            if "FAIL" in wb_status:
                failures += 1
            elif "WARN" in wb_status:
                warnings += 1
            print(_row(f"  -> vs workbook {wb_key}", _pct(engine_val),
                       _pct(wb_val), wb_status, "workbook-cache"))

    def _check_count(label: str, engine_val: int, ref_val: int):
        nonlocal failures
        status = _cmp_exact(engine_val, ref_val)
        if "FAIL" in status:
            failures += 1
        print(_row(label, _count(engine_val), _count(ref_val), status))

    # ========================================================================
    # SECTION 1: Asset counts
    # ========================================================================
    print(f"  {'-'*76}")
    print(f"  {'SECTION 1: ASSET COUNTS':^76}")
    print(f"  {'-'*76}")
    print(f"  {'Metric':<42} {'Engine':>18}  {'Reference':>18}  Status")
    print(f"  {'-'*76}")

    _check_count("Total assets loaded",       len(calc.assets),          REF["asset_count"])
    _check_count("Eligible assets",           len(calc.eligible_assets), REF["eligible_count"])
    _check_count("Ineligible assets",         len(calc.ineligible_assets), REF["ineligible_count"])
    _check_count("VAE-affected eligible",     len(calc.vae_affected),    REF["vae_affected_count"])

    print()

    # ========================================================================
    # SECTION 2: Pre-concentration values
    # ========================================================================
    print(f"  {'-'*76}")
    print(f"  {'SECTION 2: PRE-CONCENTRATION VALUES':^76}")
    print(f"  {'-'*76}")
    print(f"  {'Metric':<42} {'Engine':>18}  {'Reference':>18}  Status")
    print(f"  {'-'*76}")

    _check_dollar("Total portfolio par (OLB)",
                  calc.total_portfolio_par,    REF["total_portfolio_par"])
    _check_dollar("Pre-conc eligible ABV",
                  calc.total_pre_conc_eligible_value, REF["total_pre_conc_eligible_value"],
                  wb_key="wb_total_abv")
    _check_dollar("Pre-conc borrowing value",
                  calc.total_borrowing_value,  REF["total_borrowing_value"])

    print()

    # ========================================================================
    # SECTION 3: WAAR calculation
    # ========================================================================
    if w is None:
        print("  [SKIP] Waterfall not computed ? no eligible assets")
    else:
        print(f"  {'-'*76}")
        print(f"  {'SECTION 3: WAAR CALCULATION':^76}")
        print(f"  {'-'*76}")
        print(f"  {'Metric':<42} {'Engine':>18}  {'Reference':>18}  Status")
        print(f"  {'-'*76}")

        _check_rate("Raw WAAR (pre-cap)",
                    w.raw_waar,              REF["raw_waar"])
        _check_count("Discrete eligible obligors",
                     w.discrete_obligor_count, REF["discrete_obligor_count"])
        _check_rate("Applicable WAAR cap",
                    w.applicable_waar_cap,   REF["waar_cap"])
        _check_rate("Final WAAR (post-cap)",
                    w.final_waar,            REF["final_waar"], wb_key="wb_waar")

        print()

        # ====================================================================
        # SECTION 4: Concentration waterfall
        # ====================================================================
        print(f"  {'-'*76}")
        print(f"  {'SECTION 4: CONCENTRATION WATERFALL':^76}")
        print(f"  {'-'*76}")
        print(f"  {'Metric':<42} {'Engine':>18}  {'Reference':>18}  Status")
        print(f"  {'-'*76}")

        _check_dollar("Total excess concentration",
                      w.total_excess, REF["total_excess"], wb_key="wb_total_excess")
        _check_dollar("Net ABV (after haircuts)",
                      w.net_abv,      REF["net_abv"],      wb_key="wb_net_abv")

        # Per-test excess
        test_by_name = {t.limit_type: t for t in w.concentration_tests}
        obligor_test = test_by_name.get("Max Obligors")
        if obligor_test:
            _check_dollar("  Max Obligors excess",
                          obligor_test.excess, REF["excess_max_obligors"])
        industry_test = test_by_name.get("Max Largest Industry")
        if industry_test:
            _check_dollar("  Max Largest Industry excess",
                          industry_test.excess, REF["excess_max_largest_industry"])

        # Check all other tests have zero excess
        nonzero = {t.limit_type: t.excess for t in w.concentration_tests
                   if t.excess > 0 and t.limit_type not in {"Max Obligors", "Max Largest Industry"}}
        if nonzero:
            print(f"  {'  Unexpected non-zero tests':<42} {'':>18}  {'':>18}  "
                  f"{_col('FAIL', _RED)}")
            for name, excess in nonzero.items():
                print(f"      {name}: ${excess:,.0f}")
            failures += len(nonzero)
        else:
            print(f"  {'  All other tests zero excess':<42} {'':>18}  {'':>18}  "
                  f"{_col('PASS', _GREEN)}")

        print()

        # ====================================================================
        # SECTION 5: Availability
        # ====================================================================
        print(f"  {'-'*76}")
        print(f"  {'SECTION 5: AVAILABILITY':^76}")
        print(f"  {'-'*76}")
        print(f"  {'Metric':<42} {'Engine':>18}  {'Reference':>18}  Status")
        print(f"  {'-'*76}")

        _check_dollar("Test A ? facility cap",
                      av.test_a_facility,      REF["test_a_facility"], wb_key="wb_test_a")
        _check_dollar("Test B ? borrowing base (binding)",
                      av.test_b_borrowing_base, REF["test_b_borrowing_base"], wb_key="wb_test_b")
        _check_dollar("Test C ? credit enhancement",
                      av.test_c_credit_enhancement, REF["test_c_credit_enhancement"], wb_key="wb_test_c")
        _check_dollar("AVAILABILITY (MIN of A, B, C)",
                      av.availability,         REF["availability"],    wb_key="wb_availability")

        print()

    # ========================================================================
    # SECTION 6: Advance rate distribution (informational)
    # ========================================================================
    print(f"  {'-'*76}")
    print(f"  {'SECTION 6: ADVANCE RATE DISTRIBUTION (informational)':^76}")
    print(f"  {'-'*76}")

    from collections import Counter
    rate_dist: Counter = Counter()
    for a in calc.eligible_assets:
        rate_dist[a.advance_rate_label] += 1
    for label, cnt in sorted(rate_dist.items()):
        print(f"  {label:<50} {cnt:>3} loans")

    print()

    # ========================================================================
    # SECTION 7: Policy values discovered
    # ========================================================================
    print(f"  {'-'*76}")
    print(f"  {'SECTION 7: POLICY VALUES (loaded from AGENT sheet)':^76}")
    print(f"  {'-'*76}")

    p = data.policy
    print(f"  {'Advance rates':}")
    print(f"    FL large (>$20MM):     {p.rate_first_lien_large:.2%}")
    print(f"    FL mid ($10-20MM):     {p.rate_first_lien_mid:.2%}")
    print(f"    FL small (<$10MM):     {p.rate_first_lien_small:.2%}")
    print(f"    FILO low lev (<0.75x): {p.rate_filo_low_lev:.2%}")
    print(f"    FILO mid lev (0.75-1.25x): {p.rate_filo_mid_lev:.2%}")
    print(f"    FILO high lev (>1.25x): {p.rate_filo_high_lev:.2%}")
    print(f"    Second Lien:           {p.rate_second_lien:.2%}")
    print(f"  Concentration limits:")
    print(f"    Max Obligors:          {p.conc_max_obligor:.2%}")
    print(f"    Max Largest Industry:  {p.conc_max_largest_industry:.2%}")
    print(f"    Max Second Industry:   {p.conc_max_second_industry:.2%}")
    print(f"    Max Other Industries:  {p.conc_max_other_industries:.2%}")
    print(f"  WAAR diversity caps:")
    print(f"    High (>={p.threshold_high_diversity} obligors): {p.cap_high_diversity:.2%}")
    print(f"    Mid ({p.threshold_mid_diversity}-{p.threshold_high_diversity-1} obligors): {p.cap_mid_diversity:.2%}")
    print(f"    Low (<{p.threshold_mid_diversity} obligors):  {p.cap_low_diversity:.2%}")

    print()

    # ========================================================================
    # SUMMARY
    # ========================================================================
    print(f"  {'='*76}")
    print(f"  RESULT SUMMARY")
    print(f"  {'-'*76}")
    total = failures + warnings
    if failures == 0 and warnings == 0:
        verdict = _col("ALL CHECKS PASSED", _GREEN + _BOLD)
    elif failures == 0:
        verdict = _col(f"PASSED with {warnings} warning(s)", _YELLOW + _BOLD)
    else:
        verdict = _col(f"FAILED ? {failures} failure(s), {warnings} warning(s)", _RED + _BOLD)
    print(f"  {verdict}")
    print(f"  {'='*76}\n")

    return failures


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    wb_path = Path(r"C:\Users\henry.yan\Downloads\2025-02-11_BDC Borrowing_Base_v8.xlsm")
    if len(sys.argv) > 1:
        wb_path = Path(sys.argv[1])

    if not wb_path.exists():
        print(f"ERROR: workbook not found at {wb_path}")
        sys.exit(1)

    exit_code = run_parity(wb_path)
    sys.exit(0 if exit_code == 0 else 1)
