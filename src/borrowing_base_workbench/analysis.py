from __future__ import annotations

import json
import re
import subprocess
import warnings
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import (
    CapacityDiagnostic,
    InputField,
    OutputField,
    SheetSummary,
    ValidationRule,
    WorkbookDiagnosis,
    WritePolicy,
)

def _resolve_default_workbook() -> Path:
    candidates = [
        Path(r"C:\Users\Andres.Carrillo\OneDrive - Star Mountain Capital\03_Operations\01_Financial\Borrowing_Base\Product\2025-02-11_BDC Borrowing_Base_v8.xlsx"),
        Path(r"C:\Users\Andres.Carrillo\OneDrive - Star Mountain Capital\03_Operations\01_Financial\Borrowing_Base\Product\2025-02-11_BDC Borrowing_Base_v8.xlsm"),
        Path(r"C:\Users\Andres.Carrillo\Downloads\2025-02-11_BDC Borrowing_Base_v8.xlsx"),
        Path(r"C:\Users\Andres.Carrillo\Downloads\2025-02-11_BDC Borrowing_Base_v8.xlsm"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


DEFAULT_WORKBOOK = _resolve_default_workbook()


def _bool_flag(value: object) -> bool:
    return str(value).strip().upper() in {"Y", "YES", "TRUE"}


def _cell_str(value: object) -> str:
    return "" if value is None else str(value).strip()


def _parse_allowed_values(value: object) -> list[str]:
    text = _cell_str(value)
    if not text or text in {"None", "N/A", "Free text", "Valid date"}:
        return []
    return [part.strip() for part in text.split("|") if part.strip()]


def _sheet_summary(ws) -> SheetSummary:
    nonempty = 0
    formulas = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                nonempty += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
    return SheetSummary(ws.title, ws.max_row, ws.max_column, nonempty, formulas)


def _parse_mapping_layer_inputs(ws) -> list[InputField]:
    items: list[InputField] = []
    for row_idx in range(19, 58):
        input_id = _cell_str(ws[f"A{row_idx}"].value)
        if not input_id.startswith("IN-"):
            continue
        items.append(
            InputField(
                input_id=input_id,
                name=_cell_str(ws[f"B{row_idx}"].value),
                prompt=_cell_str(ws[f"C{row_idx}"].value),
                description=_cell_str(ws[f"D{row_idx}"].value),
                data_type=_cell_str(ws[f"E{row_idx}"].value),
                required=_cell_str(ws[f"F{row_idx}"].value) == "Y",
                allowed_values=_parse_allowed_values(ws[f"G{row_idx}"].value),
                default_rule=_cell_str(ws[f"H{row_idx}"].value),
                validation_rule=_cell_str(ws[f"I{row_idx}"].value),
                destination_sheet=_cell_str(ws[f"J{row_idx}"].value),
                destination_cell=_cell_str(ws[f"K{row_idx}"].value),
                write_method=_cell_str(ws[f"L{row_idx}"].value),
                may_infer=_bool_flag(ws[f"M{row_idx}"].value),
                human_confirm=_bool_flag(ws[f"N{row_idx}"].value),
                affects_eligibility=_cell_str(ws[f"O{row_idx}"].value),
                affects_concentration=_cell_str(ws[f"P{row_idx}"].value),
                affects_availability=_cell_str(ws[f"Q{row_idx}"].value),
                run_types=_cell_str(ws[f"R{row_idx}"].value),
                notes=_cell_str(ws[f"S{row_idx}"].value),
            )
        )
    return items


def _parse_mapping_layer_outputs(ws) -> list[OutputField]:
    items: list[OutputField] = []
    for row_idx in range(65, 93):
        output_id = _cell_str(ws[f"A{row_idx}"].value)
        if not output_id.startswith("OUT-"):
            continue
        items.append(
            OutputField(
                output_id=output_id,
                name=_cell_str(ws[f"B{row_idx}"].value),
                business_meaning=_cell_str(ws[f"C{row_idx}"].value),
                source_sheet=_cell_str(ws[f"D{row_idx}"].value),
                source_cell=_cell_str(ws[f"E{row_idx}"].value),
                output_type=_cell_str(ws[f"F{row_idx}"].value),
                required_in_response=_bool_flag(ws[f"G{row_idx}"].value),
                before_after=_cell_str(ws[f"H{row_idx}"].value),
                format_rule=_cell_str(ws[f"I{row_idx}"].value),
                capture_method=_cell_str(ws[f"J{row_idx}"].value),
                notes=_cell_str(ws[f"K{row_idx}"].value),
            )
        )
    return items


def _parse_mapping_layer_validations(ws) -> list[ValidationRule]:
    items: list[ValidationRule] = []
    for row_idx in range(100, 137):
        rule_id = _cell_str(ws[f"A{row_idx}"].value)
        if not rule_id.startswith("VR-"):
            continue
        items.append(
            ValidationRule(
                rule_id=rule_id,
                field_ref=_cell_str(ws[f"B{row_idx}"].value),
                rule_type=_cell_str(ws[f"C{row_idx}"].value),
                logic=_cell_str(ws[f"D{row_idx}"].value),
                error_message=_cell_str(ws[f"E{row_idx}"].value),
                severity=_cell_str(ws[f"F{row_idx}"].value),
                auto_correct=_bool_flag(ws[f"G{row_idx}"].value),
                escalate=_bool_flag(ws[f"H{row_idx}"].value),
                notes=_cell_str(ws[f"I{row_idx}"].value),
            )
        )
    return items


def _parse_mapping_layer_policies(ws) -> list[WritePolicy]:
    items: list[WritePolicy] = []
    for row_idx in range(143, 161):
        sheet_name = _cell_str(ws[f"A{row_idx}"].value)
        if not sheet_name or sheet_name.startswith("═"):
            continue
        items.append(
            WritePolicy(
                sheet_name=sheet_name,
                permission=_cell_str(ws[f"B{row_idx}"].value),
                allowed_range=_cell_str(ws[f"C{row_idx}"].value),
                purpose=_cell_str(ws[f"D{row_idx}"].value),
                conditions=_cell_str(ws[f"E{row_idx}"].value),
                risk_level=_cell_str(ws[f"F{row_idx}"].value),
                notes=_cell_str(ws[f"G{row_idx}"].value),
            )
        )
    return items


def _find_last_nonempty_row(ws, columns: Iterable[int] | None = None) -> int:
    columns = list(columns or range(1, ws.max_column + 1))
    for row_idx in range(ws.max_row, 0, -1):
        if any(ws.cell(row_idx, col_idx).value not in (None, "") for col_idx in columns):
            return row_idx
    return 0


def _extract_modeled_ceiling(formula: str) -> int | None:
    matches = re.findall(r"'Loan Tape - Settled'!\$[A-Z]+\$(\d+):\$[A-Z]+\$(\d+)", formula or "")
    if not matches:
        return None
    return max(int(end) for _, end in matches)


def _stage_readable_copy(workbook_path: Path) -> tuple[Path, str | None]:
    try:
        with workbook_path.open("rb") as handle:
            handle.read(1)
        return workbook_path, None
    except PermissionError:
        scratch_dir = Path(__file__).resolve().parents[2] / "scratch"
        scratch_dir.mkdir(parents=True, exist_ok=True)
        staged_path = scratch_dir / f"staged_{workbook_path.name}"
        command = [
            "powershell",
            "-NoProfile",
            "-Command",
            f"Copy-Item -LiteralPath '{workbook_path}' -Destination '{staged_path}' -Force",
        ]
        completed = subprocess.run(command, capture_output=True, text=True, check=False)
        if completed.returncode != 0:
            raise PermissionError(f"{workbook_path} could not be copied for analysis: {completed.stderr.strip()}")
        return staged_path, f"Workbook was analyzed from a staged copy because the source file is locked for direct Python reads."


def analyze_workbook(workbook_path: str | Path = DEFAULT_WORKBOOK) -> WorkbookDiagnosis:
    workbook_path = Path(workbook_path)
    readable_path, staging_note = _stage_readable_copy(workbook_path)
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="wmf image format is not supported so the image is being dropped",
            category=UserWarning,
        )
        wb = load_workbook(readable_path, data_only=False, keep_vba=True)

    mapping = wb["MAPPING_LAYER"]
    ai = wb["AI_Instructions"]
    cl = wb["Concentration Limits"]
    portfolio = wb["Portfolio"]
    loan_tape = wb["Loan Tape - Settled"]
    sm_support = wb["SM Support"]

    sheet_summaries = [_sheet_summary(ws) for ws in wb.worksheets]
    inputs = _parse_mapping_layer_inputs(mapping)
    outputs = _parse_mapping_layer_outputs(mapping)
    validations = _parse_mapping_layer_validations(mapping)
    policies = _parse_mapping_layer_policies(mapping)

    industries: list[str] = []
    for row_idx in range(81, 154):
        text = _cell_str(cl[f"E{row_idx}"].value)
        if text and "Max " not in text and text != "Industry":
            industries.append(text)

    loan_tape_last = _find_last_nonempty_row(loan_tape, columns=[7])
    sm_support_last = _find_last_nonempty_row(sm_support, columns=[2])
    portfolio_insert = 0
    for row_idx in range(1, portfolio.max_row + 1):
        if "PLACE CURSER HERE" in _cell_str(portfolio[f"A{row_idx}"].value):
            portfolio_insert = row_idx
            break

    modeled_ceiling = _extract_modeled_ceiling(_cell_str(portfolio["B49"].value))
    capacities = [
        CapacityDiagnostic(
            area="Loan Tape within current Portfolio source ranges",
            current_last_row=loan_tape_last,
            modeled_ceiling=modeled_ceiling,
            remaining_headroom=(modeled_ceiling - loan_tape_last) if modeled_ceiling else None,
            note="Portfolio formulas reference fixed Loan Tape ranges ending at row 55.",
        ),
        CapacityDiagnostic(
            area="SM Support append zone before totals",
            current_last_row=sm_support_last,
            modeled_ceiling=63,
            remaining_headroom=63 - sm_support_last,
            note="Totals begin at row 64, so scenario rows should be inserted before that summary block.",
        ),
        CapacityDiagnostic(
            area="Portfolio insertion marker",
            current_last_row=portfolio_insert,
            modeled_ceiling=None,
            remaining_headroom=None,
            note="The sheet contains a manual instruction row showing where new collateral rows were expected to be added.",
        ),
    ]

    return WorkbookDiagnosis(
        workbook_path=str(workbook_path),
        sheet_summaries=sheet_summaries,
        inputs=inputs,
        outputs=outputs,
        validation_rules=validations,
        write_policies=policies,
        capacities=capacities,
        industries=industries,
        monthly_update_tabs=[
            "Loan Tape - Settled",
            "SM Support",
            "Availability",
            "Deal Team Input",
            "Memory",
            "MAPPING_LAYER",
            "Portfolio",
        ],
        current_risks=[
            "Cached output values are mostly blank, so exact numbers need an Excel-backed calculation engine.",
            "Portfolio formulas use fixed Loan Tape ranges, which creates a monthly range-extension obligation.",
            "Deal Team Input is the UX layer, but MAPPING_LAYER is the real operational contract and should drive the app.",
            "AGENT and Availability should be admin-governed because they encode legal or facility-wide logic.",
        ],
        recommended_architecture=[
            "Use the workbook as the governed monthly logic template while the app owns workflow, validation, audit logging, and side-by-side presentation.",
            "Drive form fields, outputs, and write permissions from MAPPING_LAYER instead of hard-coding them in the app.",
            "Split the product into Deal Team mode and Admin mode. Deal Team runs scenarios; Admin manages workbook versions and monthly health checks.",
            "Treat Loan Tape - Settled, SM Support, Availability metadata, and Portfolio range health as the main monthly maintenance surfaces.",
        ],
        workbook_version=_cell_str(ai["B6"].value),
        measurement_date=_cell_str(ai["B8"].value),
        borrower_name=_cell_str(ai["B9"].value),
        facility_amount_formula=_cell_str(ai["B10"].value),
        debug={
            "portfolio_insert_row": portfolio_insert,
            "loan_tape_last_row": loan_tape_last,
            "sm_support_last_row": sm_support_last,
            "opened_workbook_path": str(readable_path),
            "staging_note": staging_note,
        },
    )


def diagnosis_to_markdown(diagnosis: WorkbookDiagnosis) -> str:
    lines = ["# Borrowing Base v8 Diagnosis", ""]
    lines.extend(f"- {line}" for line in diagnosis.summary_lines())
    lines.extend(["", "## Monthly Update Tabs", ""])
    lines.extend(f"- {item}" for item in diagnosis.monthly_update_tabs)
    lines.extend(["", "## Structural Risks", ""])
    lines.extend(f"- {item}" for item in diagnosis.current_risks)
    lines.extend(["", "## Recommended App Shape", ""])
    lines.extend(f"- {item}" for item in diagnosis.recommended_architecture)
    lines.extend(["", "## Capacity Diagnostics", ""])
    for item in diagnosis.capacities:
        lines.append(
            f"- {item.area}: last row {item.current_last_row}, modeled ceiling {item.modeled_ceiling if item.modeled_ceiling is not None else 'n/a'}, remaining headroom {item.remaining_headroom if item.remaining_headroom is not None else 'n/a'}. {item.note}"
        )
    return "\n".join(lines)


def diagnosis_to_json(diagnosis: WorkbookDiagnosis) -> str:
    payload = {
        "workbook_path": diagnosis.workbook_path,
        "summary": diagnosis.summary_lines(),
        "monthly_update_tabs": diagnosis.monthly_update_tabs,
        "current_risks": diagnosis.current_risks,
        "recommended_architecture": diagnosis.recommended_architecture,
        "capacities": [item.__dict__ for item in diagnosis.capacities],
    }
    return json.dumps(payload, indent=2)
